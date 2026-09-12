# -*- coding: utf-8 -*-
"""文档智能层集成测试：parse_document / kb_search / experience_search + chat PDF 附件。"""
import asyncio
import base64
import io
import json
import sys

import websockets

URI = "ws://127.0.0.1:9876"
results = []


def record(name, ok, detail=""):
    results.append((name, ok))
    print(f"[{'PASS' if ok else 'FAIL'}] {name}" + (f" — {detail}" if detail else ""))


async def rpc(ws, method, params=None, timeout=90):
    await ws.send(json.dumps({"jsonrpc": "2.0", "id": 1, "method": method, "params": params or {}}))
    while True:
        msg = json.loads(await asyncio.wait_for(ws.recv(), timeout=timeout))
        if msg.get("id") != 1:
            continue
        if "error" in msg:
            raise RuntimeError(f"{method}: {msg['error']}")
        return msg.get("result")


def make_test_pdf() -> bytes:
    """用 pymupdf 生成一份中文测试判决书 PDF。"""
    import pymupdf
    doc = pymupdf.open()
    page = doc.new_page()
    body = (
        "北京市海淀区人民法院\n民事判决书\n\n"
        "原告：华盛贸易有限公司，住所地上海市长宁区。\n"
        "被告：明远建设工程有限公司。\n\n"
        "本院认为：原、被告签订的《借款合同》合法有效。被告未按期归还借款本金500万元，"
        "构成违约。关于利息，双方约定年利率12%，未超过合同成立时一年期贷款市场报价利率四倍，"
        "本院予以支持。\n\n"
        "判决如下：被告于本判决生效后十日内偿还原告借款本金500万元及利息。\n\n"
        "如不服本判决，可在判决书送达之日起十五日内向本院递交上诉状。"
    )
    page.insert_text((72, 90), body, fontname="china-s", fontsize=11)
    buf = io.BytesIO()
    doc.save(buf)
    doc.close()
    return buf.getvalue()


def make_test_docx() -> bytes:
    import docx as pydocx
    d = pydocx.Document()
    d.add_heading("保密协议", level=1)
    d.add_paragraph("甲方对乙方提供的商业秘密负有保密义务，保密期限为三年。违约金为人民币伍拾万元整。")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


async def main():
    async with websockets.connect(URI, max_size=30 * 1024 * 1024) as ws:
        await asyncio.wait_for(ws.recv(), 15)  # ready

        # 1. parse_document — 中文 PDF
        pdf_b64 = base64.b64encode(make_test_pdf()).decode()
        r = await rpc(ws, "parse_document", {"filename": "判决书.pdf", "data_base64": pdf_b64})
        pdf_ok = r.get("ok") and "500万元" in r.get("text", "") and "海淀区" in r.get("text", "")
        record("parse_document PDF 中文抽取", pdf_ok,
               f"extractor={r.get('extractor')} chars={r.get('chars')}")

        # 2. parse_document — DOCX
        docx_b64 = base64.b64encode(make_test_docx()).decode()
        r = await rpc(ws, "parse_document", {"filename": "保密协议.docx", "data_base64": docx_b64})
        docx_ok = r.get("ok") and "保密" in r.get("text", "")
        record("parse_document DOCX 抽取", docx_ok, f"extractor={r.get('extractor')}")

        # 3. kb_search — 用判决书文本建库，检索违约责任
        pdf_text = (await rpc(ws, "parse_document", {"filename": "判决书.pdf", "data_base64": pdf_b64}))["text"]
        r = await rpc(ws, "kb_search", {
            "matter_id": "test-kb-matter",
            "query": "被告违约 利息 如何判决",
            "documents": [
                {"name": "判决书.pdf", "text": pdf_text},
                {"name": "保密协议.docx", "text": (await rpc(ws, "parse_document", {"filename": "保密协议.docx", "data_base64": docx_b64}))["text"]},
            ],
        })
        chunks = r.get("chunks", [])
        kb_ok = bool(chunks) and "判决书" in chunks[0].get("name", "")
        record("kb_search 命中正确文档", kb_ok,
               f"top1={chunks[0]['name'] if chunks else None} score={chunks[0].get('score') if chunks else None}")

        # 4. experience_search — 跨案件相似经验
        r = await rpc(ws, "experience_search", {
            "query": "借贷合同纠纷 违约 诉讼",
            "items": [
                {"id": "1", "matterTitle": "华盛公司", "title": "华盛与明远借贷合同纠纷", "text": "民间借贷 500万元 违约 利息 诉讼", "type": "matter"},
                {"id": "2", "matterTitle": "张某", "title": "张某诉李某离婚纠纷", "text": "离婚 抚养权 财产分割", "type": "matter"},
                {"id": "3", "matterTitle": "天行科技", "title": "发明专利侵权纠纷", "text": "专利侵权 二审 赔偿", "type": "matter"},
            ],
        })
        hits = r.get("hits", [])
        exp_ok = bool(hits) and "借贷" in hits[0].get("title", "")
        record("experience_search 相似案件召回", exp_ok,
               f"top1={hits[0].get('title') if hits else None} score={hits[0].get('score') if hits else None}")

        # 5. chat 附件：PDF 真实内容进 LLM（全链路，LLM 限流时不判失败）
        r = await rpc(ws, "chat", {
            "message": "附件判决书中判决被告偿还的金额是多少？只回答数字。",
            "files": [{"name": "判决书.pdf", "data": "data:application/pdf;base64," + pdf_b64}],
            "thread_id": "test-doc-intel",
        }, timeout=180)
        text = (r or {}).get("response", "")
        rate_limited = ("429" in text or "限流" in text)
        has_number = ("500" in text)
        record("chat PDF 附件全链路", has_number or rate_limited,
               ("限流（跳过严格断言）" if rate_limited else text[:80]))

    print("\n===== SUMMARY =====")
    fails = [n for n, ok in results if not ok]
    for n, ok in results:
        print(f"  {'PASS' if ok else 'FAIL'}  {n}")
    print(f"TOTAL: {len(results)}, FAILED: {len(fails)}")
    sys.exit(1 if fails else 0)


asyncio.run(main())
