# -*- coding: utf-8 -*-
"""doc_intel.py — 文档抽取适配层 + 案件知识库 + 经验记忆检索。

开源收口（opendataloader-pdf / RAG-Anything(MinerU) / cognee / semantica 评估结论的落地）：
- 抽取适配链（可插拔降级，装了更好的解析器自动升级）：
    PDF:  pymupdf4llm (Markdown, 最佳) → pymupdf → pdfminer.six
    DOCX: python-docx
    XLSX: openpyxl
    文本类: 直接解码
  未来接入 raganything / opendataloader 时调用 register_extractor() 即可升级首选后端。
- 案件知识库：按案件分库的 BM25 索引（jieba 分词），块级检索；
  index 以语料指纹缓存，避免每条消息重复分词。
- 经验记忆：跨案件条目（案件描述/决策/时间轴等）BM25 排序。

全部零重型依赖运行（pdfminer 兜底时 PDF 也可用），装 pymupdf 后质量自动升级。
"""
from __future__ import annotations

import base64
import hashlib
import io
import logging
import re
import threading
from typing import Callable, Optional

log = logging.getLogger("lawclaw.doc_intel")

# ─────────────────────────────────────────────────────────────
# 1. 文本抽取适配层
# ─────────────────────────────────────────────────────────────

# 外部可注册的首选后端（如 raganything / opendataloader 接入时调用）
_registered_extractor: Optional[Callable[[str, bytes], str]] = None
_register_lock = threading.Lock()


def register_extractor(fn: Callable[[str, bytes], str]) -> None:
    """注册自定义首选抽取后端。fn(filename, data_bytes) -> markdown/text。"""
    global _registered_extractor
    with _register_lock:
        _registered_extractor = fn


def _avail(name: str) -> bool:
    try:
        __import__(name)
        return True
    except Exception:
        return False


def _extract_pdf(data: bytes) -> tuple[str, str]:
    """PDF 适配链：pymupdf4llm → pymupdf → pdfminer。返回 (text, extractor)。"""
    # ① pymupdf4llm：Markdown 输出（表格/标题结构化最佳）
    try:
        import pymupdf4llm  # type: ignore
        text = pymupdf4llm.to_markdown(io.BytesIO(data))
        if text and text.strip():
            return text, "pymupdf4llm"
    except Exception as e:
        log.debug("pymupdf4llm failed: %s", e)
    # ② pymupdf 纯文本
    try:
        import pymupdf  # type: ignore
        doc = pymupdf.open(stream=data, filetype="pdf")
        pages = [page.get_text() for page in doc]
        doc.close()
        text = "\n\n".join(p for p in pages if p and p.strip())
        if text.strip():
            return text, "pymupdf"
    except Exception as e:
        log.debug("pymupdf failed: %s", e)
    # ③ pdfminer.six 兜底
    try:
        from pdfminer.high_level import extract_text  # type: ignore
        text = extract_text(io.BytesIO(data)) or ""
        if text.strip():
            return text, "pdfminer"
    except Exception as e:
        log.debug("pdfminer failed: %s", e)
    return "", "none"


def _extract_docx(data: bytes) -> tuple[str, str]:
    import docx  # python-docx
    doc = docx.Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts), "python-docx"


def _extract_xlsx(data: bytes) -> tuple[str, str]:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    lines = []
    for ws in wb.worksheets:
        lines.append(f"## 工作表: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None and str(c).strip()]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines), "openpyxl"


_TEXT_EXTS = {".md", ".txt", ".csv", ".json", ".xml", ".yaml", ".yml", ".log", ".srt"}


def extract_document_text(filename: str, data: bytes) -> dict:
    """统一抽取入口。返回 {ok, text, extractor, chars, error}。

    优先级：已注册外部后端（raganything 等）→ 按扩展名适配链。
    """
    fname = (filename or "").lower()
    # ① 已注册外部首选后端（如 raganything 的 MinerU 管线）
    if _registered_extractor is not None:
        try:
            text = _registered_extractor(fname, data)
            if text and text.strip():
                return {"ok": True, "text": text, "extractor": "registered", "chars": len(text)}
        except Exception as e:
            log.warning("registered extractor failed, falling back: %s", e)

    text, extractor, err = "", "none", None
    try:
        if fname.endswith(".pdf"):
            text, extractor = _extract_pdf(data)
        elif fname.endswith((".docx", ".doc")):
            if fname.endswith(".doc"):
                err = "旧版 .doc 请先另存为 .docx 或 PDF"
            else:
                text, extractor = _extract_docx(data)
        elif fname.endswith((".xlsx", ".xlsm")):
            text, extractor = _extract_xlsx(data)
        elif fname.endswith(".html") or fname.endswith(".htm"):
            raw = data.decode("utf-8", errors="replace")
            text = re.sub(r"<[^>]+>", " ", raw)
            text = re.sub(r"\s+", " ", text)
            extractor = "html-strip"
        elif any(fname.endswith(e) for e in _TEXT_EXTS):
            text = data.decode("utf-8", errors="replace")
            extractor = "plain"
        else:
            # 尝试按文本解码（前端白名单里的未知类型兜底）
            text = data.decode("utf-8", errors="replace")
            extractor = "plain"
            if text.count("\ufffd") > len(text) * 0.3:  # 乱码率过高 → 视为二进制
                text, extractor = "", "none"
                err = "不支持的二进制格式"
    except Exception as e:
        err = f"{type(e).__name__}: {e}"

    return {
        "ok": bool(text and text.strip()),
        "text": text or "",
        "extractor": extractor,
        "chars": len(text or ""),
        "error": err,
    }


# ─────────────────────────────────────────────────────────────
# 2. 分词与 BM25（中文：jieba，缺失时退化 bigram）
# ─────────────────────────────────────────────────────────────

def _tokenize(text: str) -> list[str]:
    text = (text or "").lower()
    try:
        import jieba  # type: ignore
        return [t for t in jieba.lcut(text) if t.strip() and t not in "，。、；：？！（）《》\"'一二三四五六七八九十"]
    except Exception:
        # bigram 兜底
        return [text[i:i + 2] for i in range(max(0, len(text) - 1))]


def chunk_text(text: str, chunk_size: int = 480, overlap: int = 60) -> list[str]:
    """按段落聚合切块（中文友好），带少量重叠。"""
    text = (text or "").strip()
    if not text:
        return []
    paras = [p.strip() for p in re.split(r"\n\s*\n|(?<=。）)|(?<=！）)|(?<=？）)", text) if p.strip()]
    chunks: list[str] = []
    buf = ""
    for p in paras:
        if len(buf) + len(p) <= chunk_size:
            buf = (buf + "\n" + p).strip()
            continue
        if buf:
            chunks.append(buf)
        if len(p) > chunk_size:
            # 长段落硬切（带 overlap）
            start = 0
            while start < len(p):
                chunks.append(p[start:start + chunk_size])
                start += chunk_size - overlap
            buf = ""
        else:
            buf = p
    if buf:
        chunks.append(buf)
    return [c for c in chunks if len(c) >= 20]


class BM25Index:
    """轻量 BM25（k1=1.5, b=0.75）。语料不变时构建一次。"""

    def __init__(self, docs: list[dict], k1: float = 1.5, b: float = 0.75):
        self.k1, self.b = k1, b
        self.docs = docs  # [{id, text, meta}]
        self.doc_tfs: list[dict[str, int]] = []
        self.doc_lens: list[int] = []
        self.df: dict[str, int] = {}
        for d in docs:
            toks = _tokenize(d.get("text", ""))
            tf: dict[str, int] = {}
            for t in toks:
                tf[t] = tf.get(t, 0) + 1
            self.doc_tfs.append(tf)
            self.doc_lens.append(len(toks))
            for t in tf:
                self.df[t] = self.df.get(t, 0) + 1
        self.avg_len = (sum(self.doc_lens) / len(self.doc_lens)) if self.doc_lens else 0.0
        self.n = len(docs)

    def search(self, query: str, top_k: int = 5) -> list[dict]:
        import math
        if not self.n or not (query or "").strip():
            return []
        q_toks = [t for t in _tokenize(query) if t in self.df]
        if not q_toks:
            return []
        scores = []
        for i in range(self.n):
            score = 0.0
            dl = self.doc_lens[i] or 1
            for t in set(q_toks):
                tf = self.doc_tfs[i].get(t, 0)
                if not tf:
                    continue
                idf = math.log(1 + (self.n - self.df[t] + 0.5) / (self.df[t] + 0.5))
                score += idf * (tf * (self.k1 + 1)) / (tf + self.k1 * (1 - self.b + self.b * dl / self.avg_len))
            if score > 0:
                scores.append((score, i))
        scores.sort(reverse=True)
        return [
            {**self.docs[i], "score": round(sc, 4)}
            for sc, i in scores[:top_k]
        ]


def _fingerprint(payload: str) -> str:
    return hashlib.sha1(payload.encode("utf-8", errors="replace")).hexdigest()[:16]


# ─────────────────────────────────────────────────────────────
# 3. 案件知识库（按案件缓存 BM25 语料）
# ─────────────────────────────────────────────────────────────

_kb_cache: dict[str, tuple[str, BM25Index]] = {}
_kb_lock = threading.Lock()
MAX_DOC_CHARS = 30000   # 单文档注入 KB 的最大字符
MAX_DOCS = 8            # 单案件最大文档数


def kb_search(matter_id: str, query: str, documents: list[dict], top_k: int = 5) -> dict:
    """案件知识库检索。documents = [{name, text}]（前端传入已抽取文本）。

    返回 {ok, chunks: [{name, text, score}], indexed_docs, indexed_chunks}。
    """
    docs_in = (documents or [])[:MAX_DOCS]
    payload = "\x00".join(d.get("text", "")[:MAX_DOC_CHARS] for d in docs_in)
    fp = _fingerprint(f"{matter_id}|{len(docs_in)}|{payload}")
    with _kb_lock:
        cached = _kb_cache.get(matter_id)
        if not cached or cached[0] != fp:
            flat: list[dict] = []
            for d in docs_in:
                name = d.get("name", "未命名")
                for i, ch in enumerate(chunk_text(d.get("text", "")[:MAX_DOC_CHARS])):
                    flat.append({"id": f"{name}#{i}", "name": name, "text": ch})
            if not flat:
                _kb_cache[matter_id] = (fp, BM25Index([]))
            else:
                _kb_cache[matter_id] = (fp, BM25Index(flat))
        index = _kb_cache[matter_id][1]
    hits = index.search(query, top_k=top_k)
    return {
        "ok": True,
        "chunks": hits,
        "indexed_docs": len(docs_in),
        "indexed_chunks": index.n,
    }


# ─────────────────────────────────────────────────────────────
# 4. 经验记忆检索（跨案件条目排序，语料由前端提供）
# ─────────────────────────────────────────────────────────────

def experience_search(query: str, items: list[dict], top_k: int = 4) -> dict:
    """经验记忆检索。items = [{id, matterTitle, type, title, text, date?}]。

    前端汇总：本案件及全部案件的描述、决策记录、时间轴事件。
    返回 {ok, hits: [{...item, score}]}。
    """
    docs = []
    for it in (items or [])[:400]:
        text = f"{it.get('matterTitle', '')} {it.get('title', '')} {it.get('text', '')}"
        if len(text.strip()) < 8:
            continue
        docs.append({**it, "text": text})
    if not docs:
        return {"ok": True, "hits": []}
    index = BM25Index(docs)
    hits = index.search(query, top_k=top_k)
    return {"ok": True, "hits": hits}
